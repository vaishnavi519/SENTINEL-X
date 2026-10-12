import io
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from pypdf import PdfReader
from docx import Document

from app.api.analyze import analyze_input
from app.config import get_settings
from app.schemas.request import AnalyzeInputRequest

router = APIRouter()

MAX_FILE_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")

        elif extension == ".pdf":
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise HTTPException(
                    status_code=422,
                    detail="Password-protected PDFs are not supported.",
                )
            text = "\n".join(page.extract_text() or "" for page in reader.pages)

        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)

        else:
            raise HTTPException(
                status_code=415,
                detail="Unsupported file type. Upload PDF, DOCX, or TXT.",
            )

    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=422,
            detail="Could not read this document. Check that it is a valid, uncorrupted file.",
        )

    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail="No extractable text found. Scanned/image-only PDFs may need OCR first.",
        )

    return text


@router.post("/analyze/document")
async def analyze_document(
    request: Request,
    file: UploadFile = File(...),
):
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file type. Upload PDF, DOCX, or TXT.",
        )

    content = await file.read(MAX_FILE_BYTES + 1)
    await file.close()

    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=413,
            detail="File is too large. Maximum upload size is 10 MB.",
        )

    text = extract_text(filename, content)
    max_chars = get_settings().max_input_chars

    if len(text) > max_chars:
        raise HTTPException(
            status_code=413,
            detail=f"Extracted text exceeds the configured limit of {max_chars} characters.",
        )

    payload = AnalyzeInputRequest(
        text=text,
        session_id=str(uuid4()),
        source="RAG_DOCUMENT",
    )

    report = await analyze_input(payload, request)

    return {
        "filename": filename,
        "extracted_characters": len(text),
        "report": report.model_dump(mode="json"),
    }
